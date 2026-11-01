"""Fill the organisers' GEH pitch template with Lodestar content. Numbers come from the app/validation."""
import sys
from pptx import Presentation
from pptx.util import Inches

TPL, IMG, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
LOGO = "docs/img/logo.png"
URL = "lodestar-utfz.onrender.com"

FILL = {
 0: [
  ("[Team name]", "Lodestar"),
  ("[Solution name]: [one sentence on what it does and for whom]",
   "Bankable demand forecasts that tell mini-grid developers and their lenders how much power a village will use, before anyone builds."),
  ("[#] [Brief title]", "03 · Predict mini-grid demand before building"),
  ("[Challenge track]", "2 · Digital Solutions & AI"),
  ("[Names and roles]", "Tushar Agarwal · product, data and engineering"),
 ],
 1: [
  ("[Who is affected, and what changes for them]", "Every mini-grid is sized on a guess about demand"),
  ("[Name the people in your brief who face this problem and how many there are]",
   "490 million people need the 217,000 mini-grids still to be built (ESMAP). Each is sized before anyone connects. Too big, and tariffs price villagers out. Too small, and evenings go dark."),
  ("\"[A real quote or data point from the community]\"",
   "\"Survey-based estimates miss individual customers' demand by 65%.\" (CrossBoundary Innovation Lab)"),
  ("[Number]", "5 of 6"),
  ("[Your brief's impact metric, e.g., households reached]",
   "years where real metered demand landed inside our forecast range (DR Congo, forecast without the site's meters)"),
  ("[From → To]", "+37% → +6%"),
  ("[Baseline today versus with your solution]",
   "oversizing at a real site: today's desk estimate vs Lodestar's plan"),
  ("[Number]", "−48%"),
  ("[Second outcome, e.g., cost saved or hours of power]",
   "cost of power when built for likely demand, not the optimistic case (20 villages)"),
 ],
 2: [
  ("[What is new about your approach]", "We forecast demand as a range lenders can bank on"),
  ("[How the problem is handled now, and why existing solutions fall short]",
   "One day-one number from a survey or benchmark: no range, no growth, never checked against meters."),
  ("[What you do differently, in one or two sentences]",
   "A probabilistic range, month by month: a safe level demand beats 9 times in 10, the likely level, and the upside."),
  ("[Which root cause from your brief this solves]",
   "No trusted demand number before building, so sites are mis-sized or never financed."),
  ("[Step 1: what happens]", "Click a village: satellite buildings and OpenStreetMap count customers"),
  ("[Step 2: what happens]", "20,000 growth scenarios give a five-year range, in plain words"),
  ("[Step 3: what happens]", "Size the system, stress-test the loan; operators' meters sharpen it"),
 ],
 3: [
  ("[How it stays affordable and pays for itself]", "Villagers pay nothing; developers and funders do"),
  ("[Cost / month]", "$0"),
  ("[Compared with what they pay today, e.g., diesel, estimated bills]",
   "for households. A developer pays $300 per site report: 0.05% of a ~$585k mini-grid."),
  ("PAYBACK", "BREAK-EVEN"),
  ("[Months or years]", "6–8 customers"),
  ("[Users, utility, government, investors: who pays what, and when]",
   "Developers: $300 per site report. Funders (RBF programmes, DFIs): $12k/yr to screen portfolios. Lenders: later, per loan file."),
  ("[Revenue, savings, or cost recovery that covers running costs without ongoing grants]",
   "Costs ~$60k/yr (two analysts, one server). Covered by 3 funder subscriptions + ~80 reports, e.g. 3–5 developers."),
  ("WHO KEEPS IT RUNNING", "MARKET"),
  ("[Local maintenance, parts, and who is responsible after launch]",
   "AMDA members plan 6,116 sites (≈$1.8M in reports). Nigeria's DARES alone: 1,000+. Needed by 2030: 217,000."),
  ("[Diesel, emissions, or waste avoided per year]",
   "No stranded panels and batteries (desk plan oversized a real site by 61 kWp, ours by 9) and less backup diesel."),
 ],
 4: [
  ("[Proof it works, and how it grows]", "Live and tested; next, a 10-site pilot in Nigeria"),
  ("[Data you tested on and the result]", "Real meters: demand in range 5 of 6 years"),
  ("[Standards, rules, or partners it fits]", "Exports to HOMER; open data and licences"),
  ("[Top risk and how you will handle it]", "Risk: few public meters → backtest operators first"),
  ("[Where, with whom, how many users]", "Nigeria: 10 sites with DARES developers"),
  ("[Next sites, utilities, or states]", "Score on meters; one funder; Kenya next"),
  ("[Countries or networks; how unit cost falls]", "5 countries: ~$30k/yr per country (est.), compute ~free"),
  ("[What you need next: pilot partner, data, funding amount, or mentor]",
   "2–3 operators' past meter data, and one RBF programme as pilot partner. They get back: tighter forecasts (cheaper debt), peer benchmarks, free reports."),
 ],
}

NOTES = {
 0: "Title. I'm Tushar, team Lodestar, Brief 03. Lodestar tells mini-grid developers and their lenders how much power a village will use, as a range they can bank on, before anyone builds.",
 1: "Impact. 490 million people need 217,000 mini-grids, and every one is sized on a guess: surveys miss what a customer uses by 65%. Too big and tariffs price villagers out; too small and evenings go dark. On two real metered mini-grids in DR Congo, forecast without their meters, real demand landed inside our range in 5 of 6 years. A desk estimate oversized one site by 37%; our plan by 6%. Building for likely demand instead of the optimistic case roughly halves the cost of power.",
 2: "Innovation. Today a developer gets one number. We give a probabilistic range, as solar finance does for sunshine: a safe level demand beats 9 times in 10, the likely level, and the upside, month by month for five years. Click a village; satellite buildings and OpenStreetMap count customers; 20,000 scenarios give the range; then size the system and stress-test the loan. Operators' meters sharpen it.",
 3: "Sustainability. Villagers never pay. Developers pay $300 per site report, 0.05% of a mini-grid; funders pay $12k a year to screen portfolios; lenders follow later. Running costs are about $60k a year, so 3 funder subscriptions and about 80 reports break even: 6 to 8 customers. The market: 6,116 sites planned by AMDA members, 1,000+ in Nigeria's DARES alone, 217,000 needed by 2030. Right-sizing also avoids stranded panels and diesel.",
 4: "Feasibility and scale. Lodestar is live, tested on real meters, and exports to HOMER. The risk is few public meters, so we test on operators' data first. Pilot: 10 sites with DARES developers in Nigeria, then one funder and Kenya, then 5 countries at about $30k a year each. Our ask: 2 to 3 operators' meter data and one RBF programme. In return they get tighter forecasts, which means cheaper debt, peer benchmarks and free reports.",
}


# The template's "Judges look for:" bar is guidance for teams; we use the same bar for our own takeaway.
BOTTOM_LINE = {
 1: "Tested on real meters: demand landed in our range 5 of 6 times, and right-sizing roughly halves the cost of power.",
 2: "A range instead of a guess, with growth over time and a lender view, tested on real meters.",
 3: "Villagers pay nothing. 6–8 paying customers cover costs, in a pipeline of thousands of planned sites.",
 4: "Live at " + URL + ". Partners who share meter data get tighter forecasts and cheaper debt in return.",
}


def set_par(par, text):
    runs = par.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def all_text_shapes(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from all_text_shapes(s.shapes)
        elif s.has_text_frame:
            yield s


prs = Presentation(TPL)
for idx, pairs in FILL.items():
    slide = prs.slides[idx]
    queue = list(pairs)
    for shp in all_text_shapes(slide.shapes):
        for par in shp.text_frame.paragraphs:
            t = "".join(r.text for r in par.runs).strip()
            for k, (old, new) in enumerate(queue):
                if t == old:
                    set_par(par, new)
                    queue.pop(k)
                    break
    if queue:
        raise SystemExit(f"slide {idx + 1}: unmatched {queue}")
    if idx in BOTTOM_LINE:
        bar = next(s for s in slide.shapes if s.name == "Judges look for text")
        runs = bar.text_frame.paragraphs[0].runs
        runs[0].text = "Bottom line:"
        runs[1].text = "  " + BOTTOM_LINE[idx]
        for r in runs[2:]:
            r.text = ""
        bar.name = "Bottom line text"
    slide.notes_slide.notes_text_frame.text = NOTES[idx]

# slide 5: replace the image placeholder with the app screenshot (forecast vs real meters)
s5 = prs.slides[4]
box = next(s for s in s5.shapes if s.name == "Prototype image placeholder")
for s in list(s5.shapes):
    if s.name in ("Placeholder icon",) or (s.has_text_frame and s.text_frame.text.startswith("[Screenshot")):
        s._element.getparent().remove(s._element)
x, y, w, h = box.left, box.top, box.width, box.height
box._element.getparent().remove(box._element)
pic = s5.shapes.add_picture(IMG, x, y, width=w)
if pic.height > h:
    ratio = h / pic.height
    pic.width, pic.height = int(pic.width * ratio), h
pic.name = "Prototype screenshot"
pic._element.nvPicPr.cNvPr.set("descr", "Lodestar app: Mutwanga year-3 forecast with P90-P10 band and the real metered demand line")
# logo above the team name on the title slide
s1 = prs.slides[0]
title = next(s for s in s1.shapes if s.has_text_frame and s.text_frame.text.strip() == "Lodestar")
lg = s1.shapes.add_picture(LOGO, title.left, title.top - Inches(0.95), width=Inches(0.8))
lg.name = "Lodestar logo"
lg._element.nvPicPr.cNvPr.set("descr", "Lodestar logo: a guiding star over a forecast band")
cp = prs.core_properties  # template metadata would otherwise label the PDF as the blank template
cp.title, cp.author, cp.last_modified_by = "Lodestar: Brief 03 Pitch, Global Energy Hackathon 2026", "Tushar Agarwal", "Tushar Agarwal"
cp.subject, cp.keywords = "Bankable demand forecasts for mini-grids", "Lodestar, mini-grid, demand forecast, P90"
prs.save(OUT)
print("saved", OUT)
